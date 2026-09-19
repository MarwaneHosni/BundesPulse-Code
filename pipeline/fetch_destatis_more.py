"""Fetch more official Destatis / VGR-Laender data and stage it.

Integrates three real, per-Bundesland datasets into the snapshot build:

* **Economy**  – Bruttoinlandsprodukt (BIP) in jeweiligen Preisen per Bundesland,
  1991-2024, plus derived BIP je Einwohner. Source: Arbeitskreis VGR der Laender
  (VGRdL) yearbook published on statistikportal.de (`vgrdl_r1b1`).
* **Housing**  – Baugenehmigungen and Baufertigstellungen (Wohnungen) per
  Bundesland, 2022. Source: Destatis *Bautaetigkeit* statistical report.
* **Transport** – Strassenverkehrsunfaelle insgesamt per Bundesland, 2024-2025.
  Source: Destatis *Statistischer Bericht Verkehrsunfaelle* (20807).

Staging CSVs are written to ``data/processed/destatis_more/`` in the same
schema used by ``pipeline/build_data.py``.
"""

from __future__ import annotations

import csv
import os
import re
import sys
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

from pipeline.source_regions import COUNTRY_ID, LAND_AGS

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

REPO_ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = REPO_ROOT / "data" / "raw" / "sources"
OUT_DIR = Path(
    os.environ.get(
        "BUNDESPULSE_PROCESSED_DIR", REPO_ROOT / "data" / "processed" / "destatis_more"
    )
)

VGRDL_URL = "https://www.statistikportal.de/sites/default/files/2026-03/vgrdl_r1b1_bs2025_0.xlsx"
VGRDL_FILE = "vgrdl_r1b1_bs2025_0.xlsx"

BAUTAETIGKEIT_URL = (
    "https://www.destatis.de/DE/Themen/Branchen-Unternehmen/Bauen/Publikationen/"
    "Downloads-Bautaetigkeit/bautaetigkeit-2050100227005.xlsx?__blob=publicationFile&v=2"
)
BAUTAETIGKEIT_FILE = "bautaetigkeit.xlsx"

VERKEHR_URLS = {
    2024: "https://www.destatis.de/DE/Themen/Gesellschaft-Umwelt/Verkehrsunfaelle/Publikationen/Downloads-Verkehrsunfaelle/statistischer-bericht-verkehrsunfaelle-jahr-2080700247005.xlsx?__blob=publicationFile&v=1",
    2025: "https://www.destatis.de/DE/Themen/Gesellschaft-Umwelt/Verkehrsunfaelle/Publikationen/Downloads-Verkehrsunfaelle/statistischer-bericht-verkehrsunfaelle-jahr-2080700257005.xlsx?__blob=publicationFile&v=1",
}
VERKEHR_FILES = {
    2024: "verkehrsunfaelle_2024.xlsx",
    2025: "verkehrsunfaelle_2025.xlsx",
}

# global indicator ids: 1-11 taken (destatis/ba/netz), 12+ new.
ID_GDP, ID_GDP_PC, ID_HP, ID_HC, ID_TRAFFIC = 12, 13, 14, 15, 16
ID_GDP_GROWTH = 27
ID_TRAFFIC_PERSON, ID_TRAFFIC_FATALITIES = 28, 29

HEADERS = {"User-Agent": "Mozilla/5.0"}
TOTAL_SITUATION = "Innerhalb und außerhalb von Ortschaften"

# csv-46241-01 columns -> indicator
VERKEHR_MEASURES = {
    "accidents": "Unfaelle_insgesamt",
    "person_accidents": "Unfaelle_Personenschaden",
    "fatalities": "Getoetete",
}


def download(url: str, dest: Path) -> None:
    if not dest.exists():
        req = urllib.request.Request(url, headers=HEADERS)
        with urllib.request.urlopen(req, timeout=180) as resp, open(dest, "wb") as out:
            out.write(resp.read())
        print(f"  downloaded {dest.name} ({dest.stat().st_size:,} bytes)")
    else:
        print(f"  using cached {dest.name}")


def _clean(v) -> float | None:
    try:
        return float(str(v).replace(" ", "").replace(",", "."))
    except (TypeError, ValueError):
        return None


def _land_id(name: str) -> str | None:
    """Map a (possibly merged-header) Bundesland name to its region id."""
    key = name.replace("_x000D_", "").replace("\n", "").replace(" ", "").lower()
    for cand, rid in LAND_AGS.items():
        if cand.lower() in key or key in cand.lower():
            return rid
    if "deutschland" in key or "bundesrepublik" in key:
        return COUNTRY_ID
    return None


# ------------------------------------------------------------------ economy (VGRdL)


def _vgr_portrait(wb) -> dict[int, dict[str, float]]:
    """Parse the FIRST block of sheet '1.1' (BIP Mio EUR) and '13.' (population)."""

    def parse_sheet(name: str) -> dict[int, dict[str, float]]:
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        # rebuilt column headers from rows 2..4
        col_id = {}
        width = max((len(r) for r in rows), default=0)
        for j in range(1, width):
            parts = [
                str(rows[i][j])
                for i in range(2, 5)
                if rows[i][j] is not None and str(rows[i][j]).strip()
            ]
            rid = _land_id("".join(parts))
            if rid:
                col_id[j] = rid
        out: dict[int, dict[str, float]] = {}
        started = False
        for r in rows:
            if r is None or all(c is None for c in r):
                if started:
                    break  # end of the first data block
                continue
            year = re.fullmatch(r"\s*(\d{4})\s*", str(r[0]).strip())
            if not year:
                if started:
                    break  # a heading row after data -> ignore further stacked blocks
                continue
            started = True
            y = int(year.group(1))
            out[y] = {}
            for j, rid in col_id.items():
                if j < len(r):
                    fv = _clean(r[j])
                    if fv is not None:
                        out[y][rid] = fv
        return out

    return parse_sheet("1.1"), parse_sheet("13.")


# ------------------------------------------------------------------ housing


def _bautaetigkeit_housing(wb) -> tuple[dict[str, float], dict[str, float]]:
    """Parse Baugenehmigungen (1.4-*) and Baufertigstellungen (1.5-*) per Land."""
    out = {"permits": {}, "completions": {}}

    def parse_sheet_group(prefix: str) -> dict[str, float]:
        values: dict[str, float] = {}
        for name in wb.sheetnames:
            if not name.startswith(prefix):
                continue
            ws = wb[name]
            rows = list(ws.iter_rows(values_only=True))
            header = wohn_row = None
            for i, r in enumerate(rows):
                unit = str(r[1] or "").strip() if len(r) > 1 else ""
                if header is None and (unit == "Einheit" or "Einheit" in unit):
                    header = i
                if (
                    wohn_row is None
                    and i > 0
                    and str(r[0] or "").strip().startswith("Wohnungen")
                ):
                    wohn_row = i
                if header is not None and wohn_row is not None:
                    break
            if header is None or wohn_row is None:
                continue
            hdr = rows[header]
            for j in range(2, len(hdr)):
                rid = _land_id(str(hdr[j]) if hdr[j] is not None else "")
                if rid and wohn_row < len(rows) and j < len(rows[wohn_row]):
                    fv = _clean(rows[wohn_row][j])
                    if fv is not None:
                        values[rid] = values.get(rid, 0.0) + fv
        return values

    out["permits"] = parse_sheet_group("1.4-")
    out["completions"] = parse_sheet_group("1.5-")
    return out["permits"], out["completions"]


# ------------------------------------------------------------------ transport


def _verkehr_measures(path: Path) -> dict[str, dict[int, dict[str, float]]]:
    """Parse per-Land road-accident measures from the 'csv-*' sheets.

    Returns ``{measure: {year: {region_id: value}}}`` for the measures in
    ``VERKEHR_MEASURES`` (total, person-injury accidents, fatalities). Only the
    grand-total 'Ortslage' row per Land is used.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    header = None
    for name in wb.sheetnames:
        if not name.lower().startswith("csv"):
            continue
        ws = wb[name]
        first = next(ws.iter_rows(values_only=True), None)
        if first and "Gebiet" in first and "Unfaelle_insgesamt" in first:
            header = first
            break
    if header is None:
        wb.close()
        return {}

    region_col = "Land" if "Land" in header else "Gebiet"
    gebiet_i = header.index(region_col)
    jahr_i = header.index("Jahr")
    ort_i = header.index("Ortslage")
    col_idx = {
        key: header.index(col) for key, col in VERKEHR_MEASURES.items() if col in header
    }
    out: dict[str, dict[int, dict[str, float]]] = {k: {} for k in col_idx}
    for name in wb.sheetnames:
        if not name.lower().startswith("csv"):
            continue
        ws = wb[name]
        for r in ws.iter_rows(values_only=True):
            if not r or len(r) <= max(gebiet_i, jahr_i, ort_i, *col_idx.values()):
                continue
            if region_col == "Gebiet" and str(r[gebiet_i]).strip() == header[gebiet_i]:
                continue  # header row
            rid = _land_id(str(r[gebiet_i]))
            if rid is None:
                continue
            if str(r[ort_i]).strip() != TOTAL_SITUATION:
                continue
            yr = re.fullmatch(r"(\d{4})", str(r[jahr_i]).strip())
            if not yr:
                continue
            for key, ci in col_idx.items():
                fv = _clean(r[ci])
                if fv is not None:
                    out[key].setdefault(int(yr.group(1)), {})[rid] = fv
    wb.close()
    return out


# ------------------------------------------------------------------ main


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    today = datetime.now(timezone.utc).date().isoformat()

    # economy
    print("- Economy: VGRdL BIP je Bundesland (1991-2024)")
    vg_file = RAW_DIR / VGRDL_FILE
    download(VGRDL_URL, vg_file)
    wb = openpyxl.load_workbook(vg_file, read_only=True, data_only=True)
    gdp_mio, pop_thousand = _vgr_portrait(wb)
    wb.close()
    print(
        f"  years: {min(gdp_mio)}-{max(gdp_mio)}, population years: {len(pop_thousand)}"
    )

    # housing
    print("- Housing: Baugenehmigungen/-fertigstellungen 2022 nach Laendern")
    bat_file = RAW_DIR / BAUTAETIGKEIT_FILE
    download(BAUTAETIGKEIT_URL, bat_file)
    wb = openpyxl.load_workbook(bat_file, read_only=True, data_only=True)
    permits, completions = _bautaetigkeit_housing(wb)
    wb.close()
    print(f"  permits: {len(permits)} Laender, completions: {len(completions)} Laender")

    # transport
    print("- Transport: Strassenverkehrsunfaelle nach Laendern (2024-2025)")
    measures: dict[str, dict[int, dict[str, float]]] = {}
    for year, url in VERKEHR_URLS.items():
        f = RAW_DIR / VERKEHR_FILES[year]
        download(url, f)
        for key, per_year in _verkehr_measures(f).items():
            measures.setdefault(key, {}).update(per_year)
    accidents = measures.get("accidents", {})
    person_accidents = measures.get("person_accidents", {})
    fatalities = measures.get("fatalities", {})
    print(f"  years: {sorted(accidents)} (person={len(person_accidents)}, fatal={len(fatalities)})")

    # derived economy: year-over-year GDP growth per Land
    gdp_growth: dict[int, dict[str, float]] = {}
    previous: dict[str, float] = {}
    for year in sorted(gdp_mio):
        for rid, v in gdp_mio[year].items():
            prev = previous.get(rid)
            if prev:
                gdp_growth.setdefault(year, {})[rid] = round((v - prev) / prev * 100.0, 2)
            previous[rid] = v

    indicators = [
        (
            ID_GDP,
            "gdp_mio",
            "Bruttoinlandsprodukt (jeweilige Preise)",
            "Economy",
            "Mio EUR",
            "BIP in jeweiligen Preisen je Bundesland (VGR der Länder)",
            "raw",
        ),
        (
            ID_GDP_PC,
            "gdp_pc",
            "Bruttoinlandsprodukt je Einwohner",
            "Economy",
            "EUR",
            "BIP je Einwohner (VGR-Werte), abgeleitet aus BIP und Einwohnern",
            "derived",
        ),
        (
            ID_HP,
            "housing_permits",
            "Baugenehmigungen (Wohnungen)",
            "Housing",
            "dwellings",
            "Genehmigte Wohnungen (Wohn- und Nichtwohngebäude), 2022",
            "raw",
        ),
        (
            ID_HC,
            "housing_completions",
            "Baufertigstellungen (Wohnungen)",
            "Housing",
            "dwellings",
            "Fertiggestellte Wohnungen (Wohn- und Nichtwohngebäude), 2022",
            "raw",
        ),
        (
            ID_TRAFFIC,
            "traffic_accidents",
            "Straßenverkehrsunfälle insgesamt",
            "Mobility",
            "accidents",
            "Straßenverkehrsunfälle insgesamt je Bundesland (Destatis 20807)",
            "raw",
        ),
        (
            ID_GDP_GROWTH,
            "gdp_growth",
            "BIP-Wachstum gegenüber Vorjahr",
            "Economy",
            "percent",
            "Jährliche Veränderung des BIP in jeweiligen Preisen (VGR der Länder)",
            "derived",
        ),
        (
            ID_TRAFFIC_PERSON,
            "traffic_person_accidents",
            "Straßenverkehrsunfälle mit Personenschaden",
            "Mobility",
            "accidents",
            "Unfälle mit Personenschaden je Bundesland (Destatis 20807)",
            "raw",
        ),
        (
            ID_TRAFFIC_FATALITIES,
            "traffic_fatalities",
            "Getötete bei Straßenverkehrsunfällen",
            "Mobility",
            "persons",
            "Bei Straßenverkehrsunfällen getötete Personen je Bundesland (Destatis 20807)",
            "raw",
        ),
    ]
    sources = [
        (
            1,
            "Statistische Ämter (Arbeitskreis VGR der Länder)",
            "BIP in jeweiligen Preisen – Länder (VGRdL r1b1, 1991-2024)",
            VGRDL_URL,
            today,
        ),
        (
            2,
            "Statistisches Bundesamt (Destatis)",
            "Bautätigkeit 2022 – Baugenehmigungen/-fertigstellungen nach Ländern",
            BAUTAETIGKEIT_URL,
            today,
        ),
        (
            3,
            "Statistisches Bundesamt (Destatis)",
            "Statistischer Bericht Straßenverkehrsunfälle (20807), 2024-2025",
            VERKEHR_URLS[2024],
            today,
        ),
    ]

    obs = []
    for year, vals in gdp_mio.items():
        for rid, v in vals.items():
            obs.append((rid, ID_GDP, year, v, 1))
            pop = pop_thousand.get(year, {}).get(rid)
            if pop:
                obs.append((rid, ID_GDP_PC, year, round(v / pop * 1000.0, 2), 1))
    for rid, v in permits.items():
        obs.append((rid, ID_HP, 2022, v, 2))
    for rid, v in completions.items():
        obs.append((rid, ID_HC, 2022, v, 2))
    for year, vals in gdp_growth.items():
        for rid, v in vals.items():
            obs.append((rid, ID_GDP_GROWTH, year, v, 1))
    for year, vals in accidents.items():
        for rid, v in vals.items():
            obs.append((rid, ID_TRAFFIC, year, v, 3))
    for year, vals in person_accidents.items():
        for rid, v in vals.items():
            obs.append((rid, ID_TRAFFIC_PERSON, year, v, 3))
    for year, vals in fatalities.items():
        for rid, v in vals.items():
            obs.append((rid, ID_TRAFFIC_FATALITIES, year, v, 3))

    _write_csv(
        OUT_DIR / "regions.csv", ["region_id", "name", "type", "parent_id", "area"], []
    )
    _write_csv(
        OUT_DIR / "indicators.csv",
        [
            "indicator_id",
            "slug",
            "name",
            "category",
            "unit",
            "description",
            "raw_or_derived",
        ],
        indicators,
    )
    _write_csv(
        OUT_DIR / "observations.csv",
        ["region_id", "indicator_id", "period", "value", "source_id"],
        obs,
    )
    _write_csv(
        OUT_DIR / "sources.csv",
        ["source_id", "provider", "dataset", "url", "retrieval_date"],
        sources,
    )

    print(f"\nwrote staging CSVs to {OUT_DIR}")
    print(f"  indicators: {len(indicators)} | observations: {len(obs)}")


def _write_csv(path: Path, header: list[str], rows: list[tuple]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        w.writerows(rows)


if __name__ == "__main__":
    main()
